#!/usr/bin/env python3
"""Batch-convert a vault's raw/ tree to markdown in raw_md/.

Usage:
    python scripts/convert_to_md.py <vault_root> [--force] [--config PATH]

Routing (one converter per extension, no cross-library retry):
    .pdf .docx .pptx .xlsx .xls -> docling-serve HTTP API (see docling_convert.py)
    .html .htm .mht .mhtml -> pandoc  (pandoc -f html -t gfm --wrap=none)
    .txt .csv            -> markitdown (required - fail if missing)
    .msg                 -> extract_msg
    .eml                 -> stdlib email
    .vsdx                -> vsdx (required - fail if missing: pip install vsdx)
                          Headless, text-only: shapes + connectors -> markdown
                          tables + Mermaid flowchart, no image rendering,
                          no OneNote/COM, LLM without vision can understand.
                          Only .vsdx; legacy binary .vsd is unsupported.
    .one .onepkg .onetoc2 -> OfficeIMO.OneNote offline parser (see onenote_conversion.py)
                          + docling/pandoc/markitdown for embedded attachments
                          Requires .NET SDK 8.0+ (once) to build scripts/OneNoteOffline;
                          published output is self-contained. No OneNote/COM needed.
    everything else      -> skipped (incl. vsd per spec, reported not retried)

Two passes: textual formats first (they feed the persistent Hebrew
dictionary), then PDFs (OCR output is checked against that dictionary but
never feeds it). Hebrew reversal fixing lives in hebrew_fix.py.

Config: <vault>/convert_config.json (see DEFAULT_CONFIG for all keys).
"""
from __future__ import annotations

import argparse
import email
import email.utils
import hashlib
import json
import re
import shutil
import sys
import tempfile
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timedelta, timezone
from pathlib import Path

import yaml

from typing import TypedDict

sys.path.insert(0, str(Path(__file__).resolve().parent))
import docling_convert
import hebrew_fix
import onenote_conversion
import vsdx_conversion

DOCLING_EXTS = {".pdf", ".docx", ".pptx", ".xlsx", ".xls"}
ONENOTE_EXTS = {".one", ".onepkg", ".onetoc2"}
VSDX_EXTS = {".vsdx"}
ROUTING = {**{e: "docling" for e in DOCLING_EXTS},
           ".txt": "markitdown", ".csv": "markitdown",
           ".msg": "msg", ".eml": "email",
           ".html": "pandoc", ".htm": "pandoc",
           ".mht": "pandoc", ".mhtml": "pandoc",
           ".vsdx": "vsdx",
           **{e: "onenote" for e in ONENOTE_EXTS}}


class DoclingCfg(TypedDict, total=False):
    url: str
    workers: int
    timeout: float
    retry_delay: float
    poll_interval: float
    health_interval: float


class PdfCfg(TypedDict):
    split_threshold: int
    chunk_pages: int


class HebrewCfg(TypedDict):
    dict_path: str
    ambiguity_margin: float


class VaultCfg(TypedDict):
    docling: DoclingCfg
    pdf: PdfCfg
    hebrew: HebrewCfg


DEFAULT_CONFIG: VaultCfg = {
    "docling": {"url": "http://localhost:5001", "workers": 1,
                "timeout": 300, "retry_delay": 1.0, "poll_interval": 2.0,
                "health_interval": 30},
    "pdf": {"split_threshold": 100, "chunk_pages": 50},
    "hebrew": {"dict_path": "data/hebrew_dict.json", "ambiguity_margin": 2.0},
}

RECENT_WINDOW = timedelta(hours=24)


def _format_down_error(exc, pending):
    header = str(exc) if exc is not None else "docling server is unreachable"
    if not pending:
        return f"{header} - no pending documents tracked"
    lines = [f"{header}", f"  {len(pending)} document(s) sent but not returned:"]
    for pp in pending[:50]:
        lines.append(f"    - {pp.as_posix()}")
    if len(pending) > 50:
        lines.append(f"    ... and {len(pending) - 50} more")
    lines.append("  The file that crashed the server is among the pending set.")
    return "\n".join(lines)


# ---------------------------------------------------------------- config

def load_config(vault_root: Path, config_path: Path | None = None) -> VaultCfg:
    cfg: VaultCfg = json.loads(json.dumps(DEFAULT_CONFIG))  # deep copy typed
    path = Path(config_path) if config_path else Path(vault_root) / "convert_config.json"
    if path.exists():
        user = json.loads(path.read_text(encoding="utf-8"))
        for section, values in user.items():
            if isinstance(values, dict) and isinstance(cfg.get(section), dict):  # type: ignore
                cfg[section].update(values)  # type: ignore
            else:
                cfg[section] = values  # type: ignore
    return cfg


# ------------------------------------------------------- frontmatter/meta

def build_frontmatter(title: str, created: datetime | None,
                      original_file: str, original_ext: str,
                      hebrew_fixed: bool,
                      original_paths: list[str] | None = None,
                      attachment_of: str | None = None) -> str:
    meta = {"title": title}
    if created is not None:
        if created.tzinfo is None:
            created = created.replace(tzinfo=timezone.utc)
        if datetime.now(timezone.utc) - created > RECENT_WINDOW:
            # Store as datetime so yaml emits a plain timestamp (no quotes around isoformat string)
            meta["created"] = created
    meta["original_file"] = original_file
    meta["original_ext"] = original_ext
    if original_paths and len(original_paths) > 1:
        meta["original_paths"] = sorted(original_paths)
    if attachment_of:
        meta["attachment_of"] = attachment_of
    if hebrew_fixed:
        meta["hebrew_fixed"] = True
    return "---\n" + yaml.safe_dump(meta, allow_unicode=True, sort_keys=False) + "---\n"


def _first_heading(md: str) -> str | None:
    for line in md.splitlines():
        m = re.match(r"^#+\s+(.+)", line)
        if m:
            return m.group(1).strip()
    return None


def _parse_pdf_date(raw: str) -> datetime | None:
    m = re.match(r"D:(\d{4})(\d{2})(\d{2})(\d{2})(\d{2})(\d{2})", raw or "")
    if not m:
        return None
    return datetime(*map(int, m.groups()), tzinfo=timezone.utc)


def extract_metadata(path: Path) -> tuple[str | None, datetime | None]:
    """Best-effort (title, created) from embedded doc metadata."""
    ext = path.suffix.lower()
    try:
        if ext == ".docx":
            import docx
            props = docx.Document(str(path)).core_properties
            return props.title or None, props.created
        if ext == ".pptx":
            import pptx
            props = pptx.Presentation(str(path)).core_properties
            return props.title or None, props.created
        if ext == ".html" or ext == ".htm":
            from bs4 import BeautifulSoup
            soup = BeautifulSoup(path.read_text(encoding="utf-8",
                                                errors="replace"), "html.parser")
            return (soup.title.string.strip() if soup.title and soup.title.string
                    else None), None
        if ext == ".pdf":
            import pypdfium2 as pdfium
            with pdfium.PdfDocument(str(path)) as pdf:
                meta = pdf.get_metadata_dict()
            return meta.get("Title") or None, _parse_pdf_date(meta.get("CreationDate"))
    except Exception:  # noqa: BLE001 - metadata is best-effort
        pass
    return None, None


def resolve_title(meta_title: str | None, md: str, path: Path) -> str:
    return meta_title or _first_heading(md) or path.stem


def resolve_created(meta_created: datetime | None, path: Path) -> datetime | None:
    if meta_created is not None:
        return meta_created
    try:
        return datetime.fromtimestamp(path.stat().st_mtime, tz=timezone.utc)
    except OSError:
        return None


# ------------------------------------------------------------ converters

def convert_txt(path: Path) -> str:
    from markitdown import MarkItDown

    return MarkItDown().convert(str(path)).text_content


def convert_html(path: Path) -> str:
    """Convert HTML/HTM via pandoc (required dependency — fails if missing)."""
    import shutil
    import subprocess

    if shutil.which("pandoc") is None:
        raise RuntimeError(
            "pandoc not found: install pandoc (https://pandoc.org/installing.html) "
            "— required for .html/.htm conversion"
        )
    result = subprocess.run(
        ["pandoc", "-f", "html", "-t", "gfm", "--wrap=none", str(path)],
        capture_output=True, text=True, encoding="utf-8",
    )
    if result.returncode != 0:
        raise RuntimeError(f"pandoc failed for {path.name}: {result.stderr.strip()[:500]}")
    return result.stdout


def _email_md(subject, sender, to, date, body) -> str:
    header = [f"Subject: {subject or ''}", f"From: {sender or ''}",
              f"To: {to or ''}", f"Date: {date or ''}"]
    return "\n".join(header) + "\n\n" + (body or "")


def convert_eml(path: Path):
    """Return (markdown, meta, attachments[(filename, bytes)])."""
    with open(path, "rb") as fh:
        msg = email.message_from_bytes(fh.read(), policy=email.policy.default)
    subject = str(msg.get("Subject") or "")
    date_raw = str(msg.get("Date") or "")
    try:
        created = email.utils.parsedate_to_datetime(date_raw) if date_raw else None
    except (TypeError, ValueError):
        created = None
    body = msg.get_body(("plain",))
    body_text = body.get_content() if body else ""
    attachments = [(a.get_filename() or "attachment", a.get_payload(decode=True) or b"")
                   for a in msg.iter_attachments()]
    md = _email_md(subject, msg.get("From"), msg.get("To"), date_raw, body_text)
    return md, (subject, created), attachments


def convert_msg(path: Path):
    """Return (markdown, meta, attachments[(filename, bytes)])."""
    import extract_msg
    msg = extract_msg.Message(str(path))
    created = None
    if msg.date:
        try:
            created = email.utils.parsedate_to_datetime(msg.date)
        except (TypeError, ValueError):
            pass
    md = _email_md(msg.subject, msg.sender, msg.to, msg.date, msg.body)
    attachments = [(a.longFilename or a.shortFilename or "attachment", a.data or b"")
                   for a in msg.attachments]
    return md, (msg.subject, created), attachments


def convert_vsdx(path: Path) -> str:
    """Convert .vsdx via vsdx_conversion (required)."""
    return vsdx_conversion.convert_vsdx(path)


def dispatch_convert(path: Path, client: docling_convert.DoclingClient,
                     cfg: VaultCfg, routing_ext: str | None = None):
    """Convert one file per the routing table.
    Returns (markdown, (meta_title, meta_created), attachments, converter_name).
    Raises on converter failure (caller records it; no cross-library retry)."""
    ext = (routing_ext if routing_ext is not None else path.suffix).lower()
    kind = ROUTING.get(ext)
    if kind is None:
        raise ValueError(f"unsupported extension: {ext}")

    # Single dispatch map replaces repeated if/switch chain (baseline smell fix)
    def _onenote():
        with tempfile.TemporaryDirectory(prefix="onenote_disp_") as tmp:
            out = Path(tmp)
            written = onenote_conversion.convert_onenote_file(path, out, path.parent, client, cfg)
            if not written:
                raise RuntimeError("onenote conversion produced no pages")
            combined = "\n\n---\n\n".join(p.read_text(encoding="utf-8") for p in written if p.suffix == ".md")
            return combined, extract_metadata(path), [], "onenote"

    def _docling():
        pdf_cfg = {"split_threshold": cfg["pdf"]["split_threshold"], "chunk_pages": cfg["pdf"]["chunk_pages"]}
        return (docling_convert.convert(path, client, pdf_cfg), extract_metadata(path), [], "docling")

    def _email():
        md, meta, atts = convert_eml(path)
        return md, meta, atts, "email"

    def _msg():
        md, meta, atts = convert_msg(path)
        return md, meta, atts, "msg"

    handlers = {
        "onenote": _onenote,
        "docling": _docling,
        "pandoc": lambda: (convert_html(path), extract_metadata(path), [], "pandoc"),
        "markitdown": lambda: (convert_txt(path), (None, None), [], "markitdown"),
        "vsdx": lambda: (convert_vsdx(path), extract_metadata(path), [], "vsdx"),
        "email": _email,
        "msg": _msg,
    }
    try:
        return handlers[kind]()  # type: ignore
    except KeyError:
        raise ValueError(f"unsupported kind: {kind}") from None


# -------------------------------------------------------------- pipeline

def should_skip(src: Path, dst: Path, force: bool) -> bool:
    if force or not dst.exists():
        return False
    return dst.stat().st_mtime >= src.stat().st_mtime


def _out_path(raw_root: Path, out_root: Path, src: Path) -> Path:
    rel = src.relative_to(raw_root)
    return (out_root / rel).with_suffix(".md")


def _file_hash(path: Path) -> str | None:
    h = hashlib.sha256()
    try:
        with open(path, "rb") as fh:
            for chunk in iter(lambda: fh.read(8192), b""):
                h.update(chunk)
    except OSError:
        return None
    return h.hexdigest()


def _resolve_canonical_link(att_hash: str, canonical: str, md_by_hash: dict[str, str], rel_by_hash: dict[str, str]) -> str:
    """Hide Message Chain for attachment canonical link resolution."""
    if att_hash in md_by_hash:
        return md_by_hash[att_hash]
    if "#" in canonical:
        for h, rel in rel_by_hash.items():
            if rel == canonical and h in md_by_hash:
                return md_by_hash[h]
        return str(Path(canonical.split("#")[0]).with_suffix(".md").as_posix())
    return str(Path(canonical.split("#")[0]).with_suffix(".md").as_posix())


def _safe_att_name(att_name: str) -> str:
    """Sanitize attachment filename to a safe .md name (M4: guard '.' / '..')."""
    base = Path(att_name).name or "attachment"
    if base in (".", ".."):
        base = "attachment"
    try:
        md_name = Path(base).with_suffix(".md").name
    except ValueError:
        md_name = "attachment.md"
    # Fallback if with_suffix produced empty or dot-only
    if not md_name or md_name in (".md",):
        md_name = "attachment.md"
    return md_name


def _patch_frontmatter_original_paths(md_path: Path, original_paths: list[str]) -> None:
    """Ensure md_path's YAML frontmatter contains original_paths (sorted)."""
    if len(original_paths) <= 1:
        return
    sorted_paths = sorted(original_paths)
    try:
        text = md_path.read_text(encoding="utf-8")
    except OSError:
        return
    if not text.startswith("---\n"):
        return
    end = text.find("\n---\n", 4)
    if end == -1:
        return
    fm_raw = text[4:end]
    body = text[end + 5 :]
    try:
        fm_data = yaml.safe_load(fm_raw) or {}
    except Exception:  # noqa: BLE001
        return
    if fm_data.get("original_paths") == sorted_paths:
        return
    fm_data["original_paths"] = sorted_paths
    # Preserve key order: ensure original_paths appears after original_ext if present.
    # yaml.safe_dump respects insertion order; move original_paths before hebrew_fixed if needed
    if "hebrew_fixed" in fm_data:
        hf = fm_data.pop("hebrew_fixed")
        fm_data["hebrew_fixed"] = hf
    new_fm = "---\n" + yaml.safe_dump(fm_data, allow_unicode=True, sort_keys=False) + "---\n"
    try:
        md_path.write_text(new_fm + body, encoding="utf-8")
    except OSError:
        pass


def _handle_onenote_file(src: Path, rel: str, out_root: Path, raw_root: Path, client, cfg: VaultCfg,
                         seen: dict[str, str], report: dict, onenote_written: list[tuple[str, Path]],
                         duplicate_groups: dict[str, list[str]] | None = None) -> bool:
    """Orchestrate one OneNote artifact; returns True if handled (report already set)."""
    file_hash = _file_hash(src)
    if file_hash is not None and file_hash in seen:
        canonical = seen[file_hash]
        report["files"][rel] = {"status": "duplicate", "duplicate_of": canonical, "hash": file_hash}
        if duplicate_groups is not None:
            duplicate_groups.setdefault(canonical, [canonical])
            if rel not in duplicate_groups[canonical]:
                duplicate_groups[canonical].append(rel)
        return True
    if file_hash is not None:
        seen[file_hash] = rel
        if duplicate_groups is not None:
            duplicate_groups.setdefault(rel, [rel])
    try:
        written = onenote_conversion.convert_onenote_file(src, out_root, raw_root, client, cfg)
        md_pages = [p for p in written if p.suffix == ".md"]
        for p in md_pages:
            onenote_written.append((rel, p))
        # Initial report: hebrew fields will be updated after dictionary + fix_text (§7/§12)
        report["files"][rel] = {"status": "converted", "converter": "onenote", "pages": len(md_pages), "hebrew_fixed": False, "ambiguous": []}
    except Exception as e:  # noqa: BLE001
        report["files"][rel] = {"status": "failed", "error": str(e)[:500]}
    return True


def run(vault_root: Path, force: bool = False,
        config_path: Path | None = None) -> dict:
    vault_root = Path(vault_root)
    cfg = load_config(vault_root, config_path)
    raw_root = vault_root / "raw"
    out_root = vault_root / "raw_md"
    out_root.mkdir(parents=True, exist_ok=True)

    client = docling_convert.DoclingClient(
        cfg["docling"]["url"], timeout=cfg["docling"]["timeout"],
        retry_delay=cfg["docling"].get("retry_delay", 1.0))
    health_interval = float(cfg["docling"].get("health_interval", 30))
    monitor = client.start_health_monitor(interval=health_interval)
    margin = cfg["hebrew"]["ambiguity_margin"]
    dict_path = vault_root / cfg["hebrew"]["dict_path"]
    _pending_before_exit = None

    all_files = sorted(p for p in raw_root.rglob("*") if p.is_file())
    report: dict = {"files": {}}

    # Notebook dirs: if a dir contains .onetoc2, its inner .one files are part of that notebook
    # and will be handled when processing the .onetoc2 itself - skip them individually.
    notebook_dirs = {p.parent for p in all_files if p.suffix.lower() == ".onetoc2"}
    onenote_inner = set()
    if notebook_dirs:
        for p in all_files:
            if p.suffix.lower() == ".one":
                for d in notebook_dirs:
                    try:
                        p.relative_to(d)
                        onenote_inner.add(p)
                        break
                    except ValueError:
                        continue

    seen: dict[str, str] = {}  # hash -> canonical rel posix
    duplicate_groups: dict[str, list[str]] = {}  # canonical rel -> all rels sharing same hash (raw files only)
    _att_md_by_hash: dict[str, str] = {}  # attachment hash -> md relative link
    _att_rel_by_hash: dict[str, str] = {}  # attachment hash -> canonical att_rel
    todo, results = [], {}
    onenote_written: list[tuple[str, Path]] = []  # (src_rel, page_path) for post-hoc Hebrew fix (§7)
    for src in all_files:
        rel = src.relative_to(raw_root).as_posix()
        ext = src.suffix.lower()
        if ext not in ROUTING:
            report["files"][rel] = {"status": "skipped",
                                    "reason": f"unsupported extension {ext}"}
            continue
        if src in onenote_inner:
            report["files"][rel] = {"status": "skipped", "reason": "part of notebook dir (handled via .onetoc2)"}
            continue
        # OneNote offline: multi-page, headless via OfficeIMO - handle immediately, not via todo batch
        if ext in ONENOTE_EXTS:
            _handle_onenote_file(src, rel, out_root, raw_root, client, cfg, seen, report, onenote_written, duplicate_groups)
            continue
        # Content-hash deduplication (before skip/force)
        file_hash = _file_hash(src)
        if file_hash is not None and file_hash in seen:
            canonical = seen[file_hash]
            report["files"][rel] = {"status": "duplicate", "duplicate_of": canonical, "hash": file_hash}
            duplicate_groups.setdefault(canonical, [canonical])
            if rel not in duplicate_groups[canonical]:
                duplicate_groups[canonical].append(rel)
            continue
        if file_hash is not None:
            seen[file_hash] = rel
            duplicate_groups.setdefault(rel, [rel])
        dst = _out_path(raw_root, out_root, src)
        if should_skip(src, dst, force):
            report["files"][rel] = {"status": "skipped", "reason": "up to date"}
            continue
        todo.append(src)

    def convert_one(src, routing_ext: str | None = None):
        if client.is_down():
            raise docling_convert.DoclingServerDown(f"docling server {client.base_url} is unreachable - aborting batch", client.pending_snapshot())
        try:
            md, meta, atts, converter = dispatch_convert(src, client, cfg, routing_ext=routing_ext)
            return src, {"md": md, "meta": meta, "attachments": atts, "converter": converter}
        except docling_convert.DoclingServerDown:
            raise
        except Exception as e:
            if client.is_down():
                raise docling_convert.DoclingServerDown(f"docling server {client.base_url} went down during conversion of {src.name}", client.pending_snapshot()) from e
            return src, {"error": str(e)}

    workers = max(1, int(cfg["docling"].get("workers", 1)))

    # Pass 1: textual formats (their converted text feeds the dictionary).
    pass1 = [s for s in todo if s.suffix.lower() != ".pdf"]
    pass2 = [s for s in todo if s.suffix.lower() == ".pdf"]

    def convert_batch(files, use_workers: bool = True):
        if not files:
            return
        if not use_workers:
            for src in files:
                src_key, res = convert_one(src)
                results[src_key] = res
        else:
            with ThreadPoolExecutor(max_workers=workers) as pool:
                for src, res in pool.map(convert_one, files):
                    results[src] = res

    try:
        convert_batch(pass1, use_workers=False)

        good_texts = [r["md"] for r in results.values() if "md" in r]
        dictionary = hebrew_fix.build_dictionary(good_texts, dict_path)

        if onenote_written:
            # OneNote is not OCR — do not apply reversal fixer (H2). Only report.
            for src_rel in {rel for rel, _ in onenote_written}:
                if src_rel in report["files"]:
                    report["files"][src_rel]["hebrew_fixed"] = False
                    report["files"][src_rel]["ambiguous"] = []

        convert_batch(pass2, use_workers=True)
    except docling_convert.DoclingServerDown as exc:
        _pending_before_exit = list(exc.pending) if exc.pending else client.pending_snapshot()
        raise docling_convert.DoclingServerDown(_format_down_error(exc, _pending_before_exit), _pending_before_exit) from exc
    finally:
        monitor.stop()
        # M2: don't swallow a real exception (e.g. RuntimeError during conversion)
        # — only raise DoclingServerDown from finally when no exception is already in flight
        if _pending_before_exit is None and client.is_down() and sys.exc_info()[0] is None:
            pending = client.pending_snapshot()
            if not pending:
                pending = [p for p in todo if p not in results]
            raise docling_convert.DoclingServerDown(_format_down_error(None, pending), pending)

    def write_output(src: Path, res: dict, dst: Path | None = None,
                     rel_key: str | None = None,
                     original_name: str | None = None,
                     original_paths: list[str] | None = None,
                     attachment_of: str | None = None,
                     parent_created: datetime | None = None):
        dst = dst or _out_path(raw_root, out_root, src)
        rel = rel_key or src.relative_to(raw_root).as_posix()
        if "error" in res:
            report["files"][rel] = {"status": "failed", "error": res["error"]}
            return None
        md = res["md"]
        # H2: dictionary-based reversal is OCR-only; final-form invariant remains for all (H1 fix)
        ocr = res.get("converter") == "docling"
        fixed, fix_report = hebrew_fix.fix_text(md, dictionary, margin=margin, ocr=ocr)
        meta_title, meta_created = res["meta"]
        # M3: attachments from NamedTemporaryFile have mtime == now; inherit parent
        if parent_created is not None:
            meta_created = parent_created
        shown = original_name or src.name
        title = resolve_title(meta_title, fixed, Path(shown))
        created = resolve_created(meta_created, src)
        fm = build_frontmatter(title, created, shown, Path(shown).suffix.lower(),
                               fix_report["hebrew_fixed"],
                               original_paths=original_paths,
                               attachment_of=attachment_of)
        dst.parent.mkdir(parents=True, exist_ok=True)
        dst.write_text(fm + "\n" + fixed, encoding="utf-8")
        report["files"][rel] = {
            "status": "converted", "converter": res["converter"],
            "hebrew_fixed": fix_report["hebrew_fixed"],
            "ambiguous": fix_report["ambiguous"]}
        return fixed

    for src in todo:
        res = results[src]
        rel = src.relative_to(raw_root).as_posix()
        op = duplicate_groups.get(rel)
        fixed = write_output(src, res, original_paths=op)
        if fixed is None:
            continue
        # Email attachments: convert into <stem>_attachments/ and link them.
        # Attachments are content-deduped against the full dataset (raw files + prior attachments).
        if res.get("attachments"):
            converted_links = []
            seen_att_dsts: set[str] = set()
            parent_md_rel = _out_path(raw_root, out_root, src).relative_to(out_root).as_posix()
            # M3: attachments inherit parent message date, not temp file mtime
            _parent_created = res.get("meta", (None, None))[1] if isinstance(res.get("meta"), tuple) else None
            for att_name, att_bytes in res["attachments"]:
                att_rel_base = f"{src.relative_to(raw_root).as_posix()}#{att_name}"
                att_rel = att_rel_base
                # Ensure unique att_rel key if duplicate attachment names in same email
                counter_rel = 1
                while att_rel in report["files"]:
                    counter_rel += 1
                    att_rel = f"{att_rel_base}_{counter_rel}"

                # Content-hash dedup against raw files and prior attachments
                att_hash = hashlib.sha256(att_bytes).hexdigest()
                if att_hash in seen:
                    canonical = seen[att_hash]
                    report["files"][att_rel] = {
                        "status": "duplicate",
                        "duplicate_of": canonical,
                        "hash": att_hash,
                    }
                    link_target = _resolve_canonical_link(att_hash, canonical, _att_md_by_hash, _att_rel_by_hash)
                    converted_links.append(f"- [{att_name}]({link_target}) (duplicate of {canonical})")
                    continue

                att_dst = (out_root / src.relative_to(raw_root).parent
                           / f"{src.stem}_attachments"
                           / _safe_att_name(att_name))
                # Handle name collisions for output path
                if att_dst.exists() or att_dst.as_posix() in seen_att_dsts:
                    stem = Path(att_name).stem or "attachment"
                    c = 1
                    candidate = att_dst
                    while candidate.exists() or candidate.as_posix() in seen_att_dsts:
                        candidate = att_dst.parent / f"{stem}_{c}.md"
                        c += 1
                    att_dst = candidate
                seen_att_dsts.add(att_dst.as_posix())
                # Use original extension for routing, not temp file suffix ambiguity
                routing_ext = Path(att_name).suffix.lower()
                with tempfile.NamedTemporaryFile(
                        suffix=Path(att_name).suffix, delete=False) as tmp:
                    tmp.write(att_bytes)
                    tmp_path = Path(tmp.name)
                try:
                    _, att_result = convert_one(tmp_path, routing_ext=routing_ext)
                    if "error" in att_result:
                        report["files"][att_rel] = {"status": "failed",
                                                     "error": att_result["error"]}
                    else:
                        write_output(tmp_path, att_result, dst=att_dst,
                                     rel_key=att_rel, original_name=att_name,
                                     attachment_of=parent_md_rel,
                                     parent_created=_parent_created)
                        # Remember this attachment's hash so later attachments dedup against it
                        if att_hash not in seen:
                            seen[att_hash] = att_rel
                            _att_rel_by_hash[att_hash] = att_rel
                        rel_link = att_dst.relative_to(out_root).as_posix()
                        _att_md_by_hash[att_hash] = rel_link
                        # Also track via canonical rel mapping for raw-file-hash dedup
                        if att_hash not in _att_rel_by_hash:
                            _att_rel_by_hash[att_hash] = att_rel
                        converted_links.append(f"- [{att_name}]({rel_link})")
                finally:
                    tmp_path.unlink(missing_ok=True)
            if converted_links:
                parent_dst = _out_path(raw_root, out_root, src)
                with open(parent_dst, "a", encoding="utf-8") as fh:
                    fh.write("\n\n## Attachments\n" + "\n".join(converted_links) + "\n")

    # Patch skipped canonicals that now have duplicates: ensure their md contains original_paths
    for canonical_rel, all_paths in duplicate_groups.items():
        if len(all_paths) <= 1:
            continue
        # canonical_rel is a raw file rel; skip synthetic att rels (# in name)
        if "#" in canonical_rel:
            continue
        canonical_src = raw_root / canonical_rel
        dst = _out_path(raw_root, out_root, canonical_src)
        if dst.exists():
            _patch_frontmatter_original_paths(dst, all_paths)

    report_path = out_root / "conversion_report.json"
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=1),
                           encoding="utf-8")

    counts = {}
    for entry in report["files"].values():
        counts[entry["status"]] = counts.get(entry["status"], 0) + 1
    print(f"converted: {counts.get('converted', 0)}  "
          f"skipped: {counts.get('skipped', 0)}  "
          f"failed: {counts.get('failed', 0)}  "
          f"duplicate: {counts.get('duplicate', 0)}  "
          f"-> {report_path}")
    # Detailed dedup log (1:1 output breaks for duplicates — see report duplicate_of)
    dups = [(rel, info["duplicate_of"], info.get("hash", "")[:8])
            for rel, info in report["files"].items() if info.get("status") == "duplicate"]
    if dups:
        print(f"dedup: {len(dups)} duplicate(s) suppressed (no raw_md output, see duplicate_of):")
        for rel, canonical, h in sorted(dups):
            print(f"  duplicate: {rel} -> {canonical}  hash:{h}")
    return report


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("vault_root", type=Path)
    ap.add_argument("--force", action="store_true",
                    help="reconvert even up-to-date outputs")
    ap.add_argument("--config", type=Path, default=None,
                    help="config file (default: <vault>/convert_config.json)")
    args = ap.parse_args(argv)
    run(args.vault_root, force=args.force, config_path=args.config)


if __name__ == "__main__":
    main()
